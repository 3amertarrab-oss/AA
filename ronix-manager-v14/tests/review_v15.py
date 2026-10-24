"""Acceptance of the seven PowerPoint review groups using synthetic data only."""
import base64,json,threading
from pathlib import Path
from functools import partial
from http.server import SimpleHTTPRequestHandler,ThreadingHTTPServer
from playwright.sync_api import sync_playwright
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'private-test-artifacts'/'review-v15';OUT.mkdir(parents=True,exist_ok=True)
checks=[]
def check(name,ok):
    assert ok,name
    checks.append(name)
    print("PASS:",name,flush=True)
class Quiet(SimpleHTTPRequestHandler):
    def log_message(self,*args):pass
server=ThreadingHTTPServer(('127.0.0.1',0),partial(Quiet,directory=str(ROOT/'dist')));threading.Thread(target=server.serve_forever,daemon=True).start()
try:
 with sync_playwright() as pw:
    browser=pw.chromium.launch(executable_path='/usr/bin/chromium',args=['--no-sandbox']);context=browser.new_context(viewport={'width':1440,'height':1100},accept_downloads=True,timezone_id='Africa/Cairo');p=context.new_page();errors=[];dialogs=[]
    p.on('pageerror',lambda e:errors.append(str(e)));p.on('dialog',lambda d:(dialogs.append(d.message),d.accept()));p.goto(f'http://127.0.0.1:{server.server_port}/RONIX_Manager.html')
    def submit():p.locator('#modalBody form').evaluate('(f)=>f.requestSubmit()')
    def term(key,name,value,kind=None):
        p.evaluate('(key)=>r15AddTerm(key)',key);row=p.locator('#'+key+' .r15-row').last;row.locator('.r15-name').fill(name);row.locator('.r15-amount').fill(str(value))
        if kind:row.locator('.r15-type').select_option(kind)
    p.evaluate('editProject()');p.locator('#pcode').fill('FX-01');p.locator('#pname').fill('مشروع تجريبي بالدولار');p.locator('#pclient').fill('عميل تجريبي');p.locator('#pcontract').fill('1000');p.locator('#pCurrency').select_option('USD');p.locator('#pRate').fill('50');p.locator('#pDate').fill('2026-10-24');term('pContractExtras','تصميم إضافي',200)
    check('project form shows original and converted supplemental total','1,200.00 USD' in p.locator('#pContractTotal').inner_text() and '60,000.00' in p.locator('#pContractTotal').inner_text())
    submit();check('currency date and supplemental contract saved',p.evaluate("data.projects[0].contract===60000&&data.projects[0].originalContract===1000&&data.projects[0].contractCurrency==='USD'&&data.projects[0].projectDate==='2026-10-24'"))
    p.evaluate("go('projects');openProject('FX-01')");check('project panel shows native-currency original extras and total',p.locator('.r15-contract-panel').count()==1 and 'تصميم إضافي' in p.locator('.r15-contract-panel').inner_text())
    p.evaluate("editAssignment('FX-01')");p.locator('#adesigner').fill('مصمم تجريبي');p.locator('#afee').fill('6000');term('aExtraWorks','رسومات إضافية',1000);term('aFeeAdjustments','زيادة متفق عليها',300,'increase');term('aFeeAdjustments','خصم اتفاق',500,'discount')
    check('designer fee preview includes work increases and discounts','6,800.00' in p.locator('#aFeeTotal').inner_text());submit()
    check('designer fee and cost center reconcile without double counting',p.evaluate("fee('FX-01')===6800&&ccAmounts('FX-01').projectCosts===6800&&ccAmounts('FX-01').projectedProfit===53200"))
    p.evaluate("editAssignment('FX-01',r14Encode('مصمم تجريبي'))");check('editing assignment retains base and detail rows',p.locator('#afee').input_value()=='6000' and p.locator('#aExtraWorks .r15-row').count()==1 and p.locator('#aFeeAdjustments .r15-row').count()==2)
    snapshot=p.evaluate('JSON.stringify(data)');p.locator('#aFeeAdjustments .r15-amount').last.fill('9000');submit();check('negative designer entitlement is rejected atomically',p.evaluate('JSON.stringify(data)')==snapshot);p.evaluate('closeModal()')
    p.evaluate("openPayment('FX-01',r14Encode('مصمم تجريبي'))");p.locator('#payamount').fill('2000');submit();check('designer payments reduce only balance not total cost',p.evaluate("paid('FX-01')===2000&&fee('FX-01')-paid('FX-01')===4800&&ccAmounts('FX-01').projectCosts===6800"))
    png=base64.b64decode(p.evaluate("()=>{const c=document.createElement('canvas');c.width=320;c.height=200;const x=c.getContext('2d');x.fillStyle='#eef2f6';x.fillRect(0,0,320,200);x.strokeStyle='#b58a48';x.lineWidth=4;x.strokeRect(60,35,200,130);x.beginPath();x.moveTo(60,90);x.lineTo(260,90);x.stroke();return c.toDataURL().split(',')[1]}"))
    p.evaluate("r14ManageProof('payment',r14Encode(data.payments[0].id))");p.locator('.r14-proof-input').set_input_files({'name':'synthetic-proof.png','mimeType':'image/png','buffer':png});p.wait_for_function("!document.querySelector('#modalBody form')._r14Proof.loading");submit()
    p.evaluate("go('payments')");p.get_by_role('button',name='حذف الدفعة',exact=True).click();p.locator('#r15DeleteReason').fill('تصحيح تسجيل مكرر');submit()
    check('deleted payment excluded from balances with proof retained',p.evaluate("data.payments.length===0&&data.deletedPayments.length===1&&paid('FX-01')===0&&Object.keys(data.paymentAttachments).length===1"))
    p.get_by_role('button',name='دفعات محذوفة (1)',exact=True).click();p.get_by_role('button',name='استعادة الدفعة',exact=True).click()
    check('restoration restores original payment and proof once',p.evaluate("data.payments.length===1&&data.deletedPayments.length===0&&paid('FX-01')===2000&&!!data.paymentAttachments[data.payments[0].id]"))

    p.evaluate("go('projects');openProject('FX-01')")
    p.get_by_role('button',name='+ دفعة مستلمة',exact=True).click()
    check('received-payment button opens selected project client',p.locator('#crClient').input_value()=='عميل تجريبي' and p.locator('.r15-selected-project .receipt-allocation').get_attribute('data-code')=='FX-01')
    p.locator('#crCurrency').select_option('USD');p.locator('#crRate').fill('51');p.locator('#crAmount').fill('400');p.locator('.receipt-allocation').fill('400');submit()
    check('partial FX receipt keeps actual cash and contractual repayment separate',p.evaluate("customerTotals('عميل تجريبي').cash===20400&&received('FX-01')===20000&&R14.fxDifference(data)===400&&R14.receivedOriginal(data,data.projects[0])===400"))
    first_id=p.evaluate('data.clientReceipts[0].id');p.evaluate("editCustomerReceipt(r14Encode(data.clientReceipts[0].batchId))")
    check('foreign receipt edits reload original currency and amounts',p.locator('#crCurrency').input_value()=='USD' and p.locator('#crAmount').input_value()=='400' and p.locator('.receipt-allocation').input_value()=='400');submit()
    check('foreign receipt edits preserve row identity and cash',p.evaluate('data.clientReceipts[0].id')==first_id and p.evaluate("customerTotals('عميل تجريبي').cash===20400"))
    snapshot=p.evaluate('JSON.stringify(data)');p.evaluate("editProject('FX-01')");p.locator('#pRate').fill('55');submit();check('historical contract rate cannot revalue existing transactions',p.evaluate('JSON.stringify(data)')==snapshot);p.evaluate('closeModal()')
    p.evaluate("openReceipt('FX-01')");p.locator('#crCurrency').select_option('USD');p.locator('#crRate').fill('52');p.locator('#crAmount').fill('800');p.locator('.receipt-allocation').fill('800');submit()
    check('full foreign settlement leaves zero native and book balance',p.evaluate("R14.receivedOriginal(data,data.projects[0])===1200&&received('FX-01')===60000&&customerTotals('عميل تجريبي').cash===62000&&R14.fxDifference(data)===2000"))
    p.evaluate("go('clients');selectClient(r14Encode('عميل تجريبي'))")
    check('client cash reconciliation includes distinct FX difference','2,000.00' in p.locator('#customerLedgerArea').inner_text())
    p.evaluate("openClaim('FX-01')");p.locator('#claimStatus').select_option('صادرة');p.evaluate("claimDraftRows=[{code:'FX-01',amount:60000}];claimRenderRows()");submit();p.evaluate('r14Match(r14Encode(data.claims[0].id))')
    matches=p.locator('.r14-match');matches.nth(0).fill('20000');matches.nth(1).fill('40000');submit()
    check('FX receipts settle the claim without counting exchange gains as debt',p.evaluate('R14.remaining(data,data.claims[0])===0&&customerTotals("عميل تجريبي").cash===62000'))

    p.evaluate("go('costcenters');showCostCenter('FX-01')");p.locator('.r15-cost-photo input').set_input_files({'name':'synthetic-project.png','mimeType':'image/png','buffer':png});p.wait_for_function('data.projects[0].photos?.length===1');p.wait_for_selector('.r15-cost-photo img')
    check('cost center upload updates shared project photo',p.locator('.r15-cost-photo img').evaluate('i=>i.complete&&i.naturalWidth>0') and p.evaluate('cover(data.projects[0]).src.startsWith("data:image/")'))
    p.locator('.r15-photo-button').click();check('cost-center image opens enlarged',p.locator('#modalBody img').is_visible());p.evaluate('closeModal()')
    p.locator('.r15-cost-photo').scroll_into_view_if_needed();p.screenshot(path=str(OUT/'cost-center.png'),full_page=True)
    p.evaluate("go('designers');showDesigner(r14Encode('مصمم تجريبي'))")
    with p.expect_download() as dl:p.get_by_role('button',name='تصدير كشف حساب المصمم Excel',exact=True).click()
    path=OUT/'designer.xlsx';dl.value.save_as(path);book=load_workbook(path,data_only=False)
    check('designer export is a real readable xlsx workbook',dl.value.suggested_filename.endswith('.xlsx') and book['ملخص المصمم']['A2'].value=='مصمم تجريبي')
    check('designer workbook contains original fees additions discounts payments and balance',list(book['المشاريع'].values)[1][3:]==(6000,1000,300,500,6800,2000,4800) and book['الدفعات']['D2'].value==2000)
    check('xlsx sheets read RTL and contain numeric money cells',book['المشاريع'].sheet_view.rightToLeft and book['المشاريع']['H2'].data_type=='n');book.close()
    with p.expect_download() as dl:p.evaluate('exportCostCenters()')
    path=OUT/'cost-centers.xlsx';dl.value.save_as(path);book=load_workbook(path)
    check('cost center workbook includes designer names without duplicating contracts',book['مراكز التكلفة']['D1'].value=='المصممون' and book['مراكز التكلفة']['D2'].value=='مصمم تجريبي' and book['مراكز التكلفة']['E2'].value==60000 and book['مراكز التكلفة'].max_row==2)
    check('workbook preserves native currency dates rates and FX difference',book['تعاقدات المشاريع']['C2'].value=='2026-10-24' and book['تعاقدات المشاريع']['F2'].value=='USD' and book['تعاقدات المشاريع']['N2'].value==2000);book.close()
    with p.expect_download() as dl:p.evaluate("exportTabs([['اختبار',[['النص','القيمة'],['=HYPERLINK(\"https://example.com\")',123.45]]]],'formula-safety')")
    path=OUT/'literal.xlsx';dl.value.save_as(path);book=load_workbook(path);check('formula-looking user text remains text in Excel',book.active['A2'].data_type=='s' and book.active['B2'].value==123.45);book.close()
    # Audit and restoration must preserve the new fields, archived payments and evidence.
    p.evaluate("go('payments');r15DeletePayment(r14Encode(data.payments[0].id))");p.locator('#r15DeleteReason').fill('اختبار سجل المحذوفات');submit()
    before=p.evaluate('data')
    with p.expect_download() as dl:p.evaluate('backup()')
    backup=Path(dl.value.path()).read_bytes();p.evaluate('data=r14Empty();r14Persist(data);render()');p.locator('#restore').set_input_files({'name':'latest.json','mimeType':'application/json','buffer':backup});p.get_by_role('button',name='استيراد هذه النسخة',exact=True).click();p.reload()
    after=p.evaluate('data');check('JSON roundtrip preserves contracts fee details currencies deleted payments and proofs',all(before[k]==after[k] for k in ['projects','assignments','payments','deletedPayments','paymentAttachments','clientReceipts','claimSettlements']))
    check('new data passes full consistency audit after restore',p.evaluate('v8Audit().length===0'))
    check('review workflows have no JavaScript exceptions',not errors)
    result={'passed':len(checks),'checks':checks,'errors':errors};(ROOT/'tests/review-v15-results.json').write_text(json.dumps(result,ensure_ascii=False,indent=2));print(json.dumps(result,ensure_ascii=False,indent=2));browser.close()
finally:server.shutdown()

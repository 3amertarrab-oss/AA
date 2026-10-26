"""Review fixes using an explicitly supplied latest backup; no private fixtures checked in."""
import argparse,json,threading
from pathlib import Path
from functools import partial
from http.server import SimpleHTTPRequestHandler,ThreadingHTTPServer
from playwright.sync_api import sync_playwright
from openpyxl import load_workbook
parser=argparse.ArgumentParser();parser.add_argument('--backup',type=Path,required=True);args=parser.parse_args()
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'private-test-artifacts/review-v17';OUT.mkdir(parents=True,exist_ok=True);checks=[]
def check(name,ok):
 assert ok,name
 checks.append(name);print('PASS:',name,flush=True)
class Quiet(SimpleHTTPRequestHandler):
 def log_message(self,*args):pass
server=ThreadingHTTPServer(('127.0.0.1',0),partial(Quiet,directory=str(ROOT/'dist')));threading.Thread(target=server.serve_forever,daemon=True).start();source=json.loads(args.backup.read_text())
try:
 with sync_playwright() as pw:
  browser=pw.chromium.launch(executable_path='/usr/bin/chromium',args=['--no-sandbox']);ctx=browser.new_context(viewport={'width':1440,'height':1080},accept_downloads=True,timezone_id='Africa/Cairo');p=ctx.new_page();errors=[];dialogs=[];p.on('pageerror',lambda e:errors.append(str(e)));p.on('dialog',lambda d:(dialogs.append(d.message),d.accept()));p.goto(f'http://127.0.0.1:{server.server_port}/RONIX_Manager.html')
  def submit():p.locator('#modalBody form').evaluate('(f)=>f.requestSubmit()')
  check('backup and restore removed from top bar',p.locator('.top button').count()==1)
  p.evaluate('go("reports")');check('reports exposes backup restore and recycle bin',p.get_by_role('button',name='↓ نسخة احتياطية JSON',exact=True).count()==1 and p.get_by_role('button',name='↑ استعادة نسخة',exact=True).count()==1);p.screenshot(path=str(OUT/'01-reports-backup.png'),full_page=True)
  p.locator('#restore').set_input_files(str(args.backup));p.get_by_role('button',name='استيراد هذه النسخة',exact=True).click()
  check('latest source records preserved without resurrecting older data',p.evaluate('(before)=>["projects","assignments","payments","claims","clientReceipts","receiptAttachments","paymentAttachments"].every(k=>JSON.stringify(data[k])===JSON.stringify(before[k]))',source))
  check('latest import passes accounting audit',p.evaluate('v8Audit().length===0'))
  # The reported 6,500 EGP / SAR 14 case, entered from cash rather than guessed native amount.
  foreign=p.evaluate('data.projects.find(p=>R14.projectCurrency(p)==="SAR"&&R14.projectRate(p)===14).code')
  p.evaluate('(code)=>{go("projects");openProject(r14Encode(code));openReceipt(r14Encode(code))}',foreign);p.locator('#crAmount').fill('6500')
  row=p.locator('.receipt-allocation[data-code="'+foreign+'"]');cash=row.locator('xpath=ancestor::tr').locator('.r15-receipt-cash')
  cash.fill('6500');check('EGP input calculates SAR repayment without changing actual cash',row.input_value()=='464.29' and cash.input_value()=='6500');p.screenshot(path=str(OUT/'02-egp-to-sar-receipt.png'),full_page=True);submit()
  check('cross-currency receipt saves exact cash and explicit rounding difference',p.evaluate('(code)=>{const r=data.clientReceipts.filter(x=>x.code===code).at(-1);return r.receiptOriginal===6500&&r.cashAmount===6500&&r.projectAmount===464.29&&r.amount===6500.06&&R14.delta(r.cashAmount,r.amount)===-.06}',foreign))
  p.evaluate('(code)=>editCustomerReceipt(r14Encode(data.clientReceipts.filter(x=>x.code===code).at(-1).batchId))',foreign);check('receipt edit reloads actual and contractual amounts separately',p.locator('.r15-receipt-cash[data-code="'+foreign+'"]').input_value()=='6500' and p.locator('.receipt-allocation[data-code="'+foreign+'"]').input_value()=='464.29');submit()
  # The user enters only actual cash; contractual currency is read-only and derived.
  p.evaluate('(code)=>openReceipt(r14Encode(code))',foreign);p.locator('#crAmount').fill('1000');native=p.locator('.receipt-allocation[data-code="'+foreign+'"]');native.locator('xpath=ancestor::tr').get_by_role('button',name='استخدام باقي السند').click();check('one cash input allocates remainder and calculates foreign contract settlement',p.locator('.r15-receipt-cash[data-code="'+foreign+'"]').input_value()=='1000' and native.input_value()=='71.43' and native.get_attribute('readonly') is not None);p.evaluate('closeModal()')
  # The first spreadsheet must expose all meaningful data, not only four header fields.
  project=p.evaluate('data.projects.find(p=>data.assignments.filter(a=>a.code===p.code).length>1).code')
  with p.expect_download() as dl:p.evaluate('(code)=>projectExcel(r14Encode(code))',project)
  path=OUT/'project-report.xlsx';dl.value.save_as(path);book=load_workbook(path);first=book.worksheets[0];values=dict((r[0],r[1]) for r in list(first.values)[1:] if r[0]);stats=p.evaluate('(code)=>R16.project(data,code)',project)
  check('first Excel sheet contains project totals and full transaction overview',first.title=='كشف المشروع' and first.max_row>25 and values['قيمة العقد بالجنيه']==stats['contract'] and values['المدفوع للمصممين']==stats['designers']['paid'] and values['القبض الفعلي']==stats['cash'])
  names=p.evaluate('(code)=>data.assignments.filter(a=>a.code===code).map(a=>a.designer)',project);check('first sheet names every assigned designer',all(n+' — الأتعاب' in values for n in names));check('detail sheets include numeric payments and links',book['المصممون'].max_row==len(names)+1 and 'الربط المحاسبي' in book.sheetnames and book['دفعات المصممين'].max_row>1);book.close()
  p.evaluate('editProject()');p.locator('#pcode').fill('TAX-01');p.locator('#pname').fill('مشروع مع ضريبة');p.locator('#pcontract').fill('10000');p.get_by_role('button',name='+ إضافة ضريبة',exact=True).click();p.screenshot(path=str(OUT/'03-project-tax.png'),full_page=True);p.locator('#overlay').click(position={'x':2,'y':2});check('backdrop click retains open form and typed data',p.locator('#overlay').evaluate('e=>e.classList.contains("open")') and p.locator('#pcode').input_value()=='TAX-01');p.keyboard.press('Escape');check('Escape closes after warning about unsaved input',not p.locator('#overlay').evaluate('e=>e.classList.contains("open")') and any('غير محفوظة' in x for x in dialogs));p.evaluate('editProject()');p.locator('#pname').fill('Draft name');p.locator('.modalhead .close').click();check('explicit close button dismisses dialog after the same warning',not p.locator('#overlay').evaluate('e=>e.classList.contains("open")'))
  # Deletion checks use a synthetic snapshot only.
  p.evaluate('data=R14.normalize({projects:[{code:"P",name:"Test project",client:"Test client",contract:1000,photos:[]}],assignments:[{id:"A",code:"P",designer:"Test designer",fee:200}],payments:[]});r14Persist(data);go("projects");openProject("P")');p.screenshot(path=str(OUT/'04-delete-project.png'),full_page=True)
  p.get_by_role('button',name='حذف المشروع',exact=True).click();p.locator('#r17DeleteReason').fill('Synthetic duplicate');submit();check('unused project and assignment move to recycle bin',p.evaluate('data.projects.length===0&&data.assignments.length===0&&data.trashRecords.length===1'))
  p.evaluate('r17Trash()');p.get_by_role('button',name='استعادة السجل',exact=True).click();check('project restore returns assignment and original balances',p.evaluate('getP("P").contract===1000&&fee("P")===200&&data.trashRecords.length===0'))
  p.evaluate('go("projects");openProject("P")');p.get_by_role('button',name='حذف التكليف',exact=True).click();p.locator('#r17DeleteReason').fill('Wrong assignment');submit();check('assignment deletion leaves project intact',p.evaluate('data.projects.length===1&&data.assignments.length===0'));p.evaluate('r17Trash()');p.get_by_role('button',name='استعادة السجل',exact=True).click()
  p.evaluate('go("designers");showDesigner(r14Encode("Test designer"))');p.screenshot(path=str(OUT/'05-delete-designer.png'),full_page=True);p.get_by_role('button',name='حذف المصمم',exact=True).click();p.locator('#r17DeleteReason').fill('Duplicate designer');submit();check('unpaid designer can be deleted and restored',p.evaluate('data.assignments.length===0&&data.trashRecords[0].kind==="designer"'));p.evaluate('r17Trash()');p.get_by_role('button',name='استعادة السجل',exact=True).click()
  p.evaluate('r14Change("fixture",d=>d.payments.push({id:"PAY",code:"P",designer:"Test designer",date:"2026-10-25",amount:50}));go("projects");openProject("P")');p.get_by_role('button',name='حذف المشروع',exact=True).click();check('linked project deletion explains accounting-preserving archive','أرشفة' in p.locator('#modalTitle').inner_text());p.locator('#r17DeleteReason').fill('Finished');submit();check('archived project retains financial balances and payment',p.evaluate('getP("P").archivedAt&&ccAmounts("P").designerCash===50&&data.payments.length===1'));check('archived project hidden from working directory',p.locator('#projectRows').get_by_role('button',name='فتح ←',exact=True).count()==0);p.locator('#r17ArchivedProjects').check();check('archived project remains accessible on demand',p.locator('#projectRows').get_by_role('button',name='فتح ←',exact=True).count()==1)
  p.evaluate('go("designers");showDesigner(r14Encode("Test designer"))');p.get_by_role('button',name='حذف المصمم',exact=True).click();p.locator('#r17DeleteReason').fill('Inactive');submit();check('paid designer archive preserves its ledger',p.evaluate('data.archivedDesigners.length===1&&paid("P","Test designer")===50&&fee("P","Test designer")===200'))
  before=p.evaluate('data')
  with p.expect_download() as dl:p.evaluate('backup()')
  backup=Path(dl.value.path()).read_bytes();p.locator('#restore').set_input_files({'name':'roundtrip.json','mimeType':'application/json','buffer':backup});p.get_by_role('button',name='استيراد هذه النسخة',exact=True).click();p.reload();check('archive state and financial links survive backup restore',all(p.evaluate('data')[k]==before[k] for k in ['projects','assignments','payments','archivedDesigners','trashRecords']))
  check('review workflows remain free of script exceptions',not errors)
  result={'passed':len(checks),'checks':checks,'errors':errors};(ROOT/'tests/review-v17-results.json').write_text(json.dumps(result,ensure_ascii=False,indent=2));print(json.dumps(result,ensure_ascii=False,indent=2));browser.close()
finally:server.shutdown()

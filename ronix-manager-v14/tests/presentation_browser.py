"""V19 real UI: stable document labels, whole-money display and optional FX lookup."""
import json,threading
from pathlib import Path
from functools import partial
from http.server import SimpleHTTPRequestHandler,ThreadingHTTPServer
from playwright.sync_api import sync_playwright
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'private-test-artifacts/presentation-v19';OUT.mkdir(parents=True,exist_ok=True);checks=[]
def check(name,ok):
 assert ok,name
 checks.append(name);print('PASS:',name,flush=True)
class Quiet(SimpleHTTPRequestHandler):
 def log_message(self,*args):pass
server=ThreadingHTTPServer(('127.0.0.1',0),partial(Quiet,directory=str(ROOT/'dist')));threading.Thread(target=server.serve_forever,daemon=True).start()
fixture={'projects':[{'code':'P','name':'مشروع تجريبي','client':'عميل تجريبي','contract':1000.97},{'code':'FX','name':'مشروع ريال','client':'عميل تجريبي','contract':1400,'originalContract':100,'contractCurrency':'SAR','contractRate':14}], 'assignments':[{'id':'D','code':'P','designer':'مصمم تجريبي','fee':100}], 'payments':[{'id':'PAY-aaaaaaaa-bbbb-cccc-dddd-eeeeeeeeeeee','code':'P','designer':'مصمم تجريبي','date':'2026-10-01','amount':10.25}], 'claims':[{'id':'CL-aaaaaaaa-bbbb-cccc-dddd-eeeeeeeeeeee','code':'P','client':'عميل تجريبي','date':'2026-10-01','status':'معتمدة','total':300.75,'matchMode':'manual'}], 'clientReceipts':[{'id':'R-aaaaaaaa-bbbb-cccc-dddd-eeeeeeeeeeee','batchId':'RC-aaaaaaaa-bbbb-cccc-dddd-eeeeeeeeeeee','code':'P','client':'عميل تجريبي','date':'2026-10-02','amount':25.25}], 'claimSettlements':[{'claimId':'CL-aaaaaaaa-bbbb-cccc-dddd-eeeeeeeeeeee','receiptId':'R-aaaaaaaa-bbbb-cccc-dddd-eeeeeeeeeeee','code':'P','amount':20.25}]}
try:
 with sync_playwright() as pw:
  # The sandbox outbound proxy uses a private certificate; this exception is test-only.
  browser=pw.chromium.launch(executable_path='/usr/bin/chromium',args=['--no-sandbox']);ctx=browser.new_context(ignore_https_errors=True,viewport={'width':1440,'height':1080},accept_downloads=True,timezone_id='Africa/Cairo');p=ctx.new_page();errors=[];dialogs=[];requests=[]
  p.on('pageerror',lambda e:errors.append(str(e)));p.on('dialog',lambda d:(dialogs.append(d.message),d.accept()));p.on('request',lambda r:requests.append(r.url) if r.url.startswith('https:') else None)
  url=f'http://127.0.0.1:{server.server_port}/RONIX_Manager.html';p.goto(url)
  def submit():p.locator('#modalBody form').evaluate('(f)=>f.requestSubmit()')
  p.evaluate('(d)=>{r14Persist(R14.normalize(d));render()}',fixture)
  source=p.evaluate('data');check('migration preserves old financial records and source IDs',all(source[k]==fixture[k] for k in ['projects','assignments','payments','clientReceipts','claimSettlements']) and source['claims'][0]['id']==fixture['claims'][0]['id'])
  check('dashboard shows rounded whole money without changing cash',p.evaluate('fmt(388361.97)==="388,362"&&fmt(-0.25)==="0"&&customerTotals("عميل تجريبي").cash===25.25'))
  p.evaluate('go("claims")');check('claim list shows a short code and no UUID', 'CL-0001' in p.locator('.r14-claims-table').inner_text() and 'aaaaaaaa-bbbb' not in p.locator('.r14-claims-table').inner_text());p.get_by_role('button',name='تعديل',exact=True).first.click();check('short claim code still opens original editable claim',p.evaluate('window.claimEditingId')==fixture['claims'][0]['id']);p.evaluate('closeModal()')
  html=p.evaluate('r14ClaimDocument(data.claims[0])');check('claim print uses short code and whole amounts with exact underlying settlement', 'CL-0001' in html and '300.75' not in html and '280.5' not in html and p.evaluate('R14.remaining(data,data.claims[0])')==280.5)
  p.evaluate('go("payments")');check('designer payment table uses short payment number','PV-0001' in p.locator('#content').inner_text() and 'PAY-aaaaaaaa' not in p.locator('#content').inner_text())
  p.evaluate('selectClient(r14Encode("عميل تجريبي"))');check('receipt is consistently numbered in customer ledger','RC-0001' in p.locator('#customerLedgerArea').inner_text() and 'RC-aaaaaaaa' not in p.locator('#customerLedgerArea').inner_text())
  numbers=p.evaluate('data.documentNumbers');p.reload();check('document numbers survive reload unchanged',p.evaluate('data.documentNumbers')==numbers)
  p.evaluate('openCustomerReceipt(r14Encode("عميل تجريبي"))');check('EGP rate remains one with lookup disabled',p.locator('#crRateUpdate').is_disabled() and p.locator('#crRate').input_value()=='1')
  check('rate form links to the official Banque Misr page',p.get_by_role('link',name='أسعار بنك مصر الرسمية').get_attribute('href').startswith('https://www.banquemisr.com/'))
  check('opening forms makes no external requests',not requests)
  today=p.evaluate('r14Today()');payload={'date':today,'sar':{'egp':14.72523683},'usd':{'egp':55.2}}
  def good(route):route.fulfill(status=200,content_type='application/json',headers={'Access-Control-Allow-Origin':'*'},body=json.dumps(payload))
  ctx.route('https://cdn.jsdelivr.net/**',good)
  p.locator('#crCurrency').select_option('SAR');p.locator('#crAmount').fill('10');before=p.evaluate('JSON.stringify(data)');p.locator('#crRateUpdate').click();p.wait_for_function('document.getElementById("crRate").value==="14.725237"')
  check('daily lookup fills selected currency rate and publication date',today in p.locator('#crRateStatus').inner_text() and 'مرجعي' in p.locator('#crRateStatus').inner_text())
  check('rate lookup alone changes no saved financial data',p.evaluate('JSON.stringify(data)')==before)
  submit();row=p.evaluate('data.clientReceipts.at(-1)');check('saved receipt retains quote provenance and exact converted cash',row['receiptRate']==14.725237 and row['rateQuote']['asOf']==today and row['receiptOriginal']==10 and row['cashAmount']==147.25)
  batch=row['batchId'];check('new receipt receives next short code with old code retained',p.evaluate('(id)=>r19Doc("receipt",id)',batch)=='RC-0002' and p.evaluate('r19Doc("receipt",data.clientReceipts[0].id)')=='RC-0001')
  p.evaluate('(id)=>editCustomerReceipt(r14Encode(id))',batch);check('editing shows saved source and reference number',today in p.locator('#crRateStatus').inner_text() and 'RC-0002' in p.locator('#modalBody').inner_text());p.locator('#crRate').fill('14.5');check('manual rate points to Banque Misr guidance','بنك مصر' in p.locator('#crRateStatus').inner_text());submit();check('manual override persists with Banque Misr provenance without changing receipt identity',p.evaluate('(id)=>{const r=data.clientReceipts.find(r=>r.batchId===id);return r.rateQuote.mode==="manual"&&r.rateQuote.source==="Banque Misr (manual)"&&r.rateQuote.asOf&&r.receiptRate===14.5&&R14.cashReceipt(r)===145&&r19Doc("receipt",id)==="RC-0002"}',batch))
  p.evaluate('openCustomerReceipt(r14Encode("عميل تجريبي"))');p.locator('#crCurrency').select_option('SAR');p.locator('#crRate').fill('13.5');ctx.unroute('https://cdn.jsdelivr.net/**',good)
  def fail(route):route.abort()
  ctx.route('https://cdn.jsdelivr.net/**',fail);ctx.route('https://latest.currency-api.pages.dev/**',fail);p.locator('#crRateUpdate').click();p.wait_for_function('document.getElementById("crRateStatus").textContent.includes("تعذر")');check('network failure retains previous rate and offers manual entry',p.locator('#crRate').input_value()=='13.5' and not p.locator('#crRateUpdate').is_disabled())
  ctx.unroute('https://latest.currency-api.pages.dev/**',fail);ctx.route('https://latest.currency-api.pages.dev/**',good);p.locator('#crRateUpdate').click();p.wait_for_function('document.getElementById("crRate").value==="14.725237"');check('fallback provider succeeds after primary failure',today in p.locator('#crRateStatus').inner_text())
  ctx.unroute('https://cdn.jsdelivr.net/**',fail);ctx.unroute('https://latest.currency-api.pages.dev/**',good)
  # Delayed responses must not overwrite a manual edit or a different currency/form.
  pending=[]
  def delay(route):pending.append(route)
  ctx.route('https://cdn.jsdelivr.net/**',delay)
  p.locator('#crRateUpdate').click();p.wait_for_timeout(100);p.locator('#crRate').fill('16.123456');good(pending.pop());p.wait_for_timeout(150);check('late response cannot overwrite manual entry',p.locator('#crRate').input_value()=='16.123456')
  p.locator('#crRateUpdate').click();p.wait_for_timeout(100);p.locator('#crCurrency').select_option('USD');p.locator('#crRate').fill('55');good(pending.pop());p.wait_for_timeout(150);check('late response for old currency is ignored',p.locator('#crRate').input_value()=='55' and p.locator('#crCurrency').input_value()=='USD')
  p.locator('#crRateUpdate').click();p.wait_for_timeout(100);p.evaluate('closeModal();openCustomerReceipt(r14Encode("عميل تجريبي"))');good(pending.pop());p.wait_for_timeout(150);check('response from a closed form cannot affect the next form',p.locator('#crRate').input_value()=='1' and p.locator('#crCurrency').input_value()=='EGP');ctx.unroute('https://cdn.jsdelivr.net/**',delay)
  ctx.route('https://cdn.jsdelivr.net/**',good);p.evaluate('closeModal();editProject(r14Encode("P"))');check('historical linked contract rate cannot be refreshed',p.locator('#pRateUpdate').is_disabled() and p.locator('#pRate').get_attribute('readonly') is not None);p.evaluate('closeModal();editProject()');p.locator('#pCurrency').select_option('SAR');p.locator('#pRateUpdate').click();p.wait_for_function('document.getElementById("pRate").value==="14.725237"');check('new project also supports daily lookup with full rate precision',p.locator('#pRate').input_value()=='14.725237');p.evaluate('closeModal()')
  with p.expect_download() as dl:p.evaluate('projectExcel(r14Encode("P"))')
  path=OUT/'project.xlsx';dl.value.save_as(path);book=load_workbook(path);overview=book['كشف المشروع'];money_cell=next(row[1] for row in overview if row[0].value=='قيمة العقد بالجنيه');check('Excel displays whole money but retains exact numeric value',money_cell.value==1000.97 and money_cell.number_format=='#,##0');book.close()
  with p.expect_download() as dl:p.evaluate('exportTabs([["Precision",[["المبلغ","سعر الصرف"],[10.25,14.725237]]]],"Precision")')
  path=OUT/'precision.xlsx';dl.value.save_as(path);book=load_workbook(path);check('Excel exchange rates retain fractional display',book.active['A2'].number_format=='#,##0' and book.active['B2'].number_format=='0.######' and book.active['B2'].value==14.725237);book.close()
  with p.expect_download() as dl:p.evaluate('backup()')
  backup=Path(dl.value.path()).read_bytes();saved=json.loads(backup);check('backup keeps exact amounts source IDs short codes and quote provenance',saved['documentNumbers']==p.evaluate('data.documentNumbers') and saved['projects'][0]['contract']==1000.97 and saved['clientReceipts'][-1]['rateQuote']['mode']=='manual')
  fresh=browser.new_context();q=fresh.new_page();q.on('dialog',lambda d:d.accept());q.on('pageerror',lambda e:errors.append(str(e)));q.goto(url);q.locator('#restore').set_input_files({'name':'V19.json','mimeType':'application/json','buffer':backup});q.get_by_role('button',name='استيراد هذه النسخة',exact=True).click();check('fresh restore preserves all numbers cash and links',q.evaluate('data.documentNumbers')==saved['documentNumbers'] and q.evaluate('data.claimSettlements')==saved['claimSettlements'] and q.evaluate('customerTotals("عميل تجريبي").cash')==170.25)
  # Actual public endpoint smoke check, separate from deterministic accounting tests.
  ctx.unroute('https://cdn.jsdelivr.net/**',good);p.evaluate('openCustomerReceipt(r14Encode("عميل تجريبي"))');p.locator('#crCurrency').select_option('SAR');p.locator('#crRateUpdate').click();p.wait_for_function('!document.getElementById("crRateUpdate").disabled',timeout=25000);check('live public daily rate fetch works from browser',p.evaluate('document.getElementById("crRate")._r19Quote?.currency')=='SAR' and float(p.locator('#crRate').input_value())>1)
  p.screenshot(path=str(OUT/'rate-picker.png'),full_page=True);p.evaluate('closeModal();go("claims")');p.wait_for_timeout(250);p.screenshot(path=str(OUT/'short-claims.png'),full_page=True)
  check('V19 workflows have no JavaScript exceptions',not errors)
  (ROOT/'tests/presentation-browser-results.json').write_text(json.dumps({'passed':len(checks),'checks':checks,'errors':errors},ensure_ascii=False,indent=2));browser.close()
finally:server.shutdown()
